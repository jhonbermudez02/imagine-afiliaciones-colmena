#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


@dataclass
class InputSet:
    afiliados_csv: Path
    devuelto_csv: Path
    dewvuelto_csv: Path
    informe_xlsx: Path
    resultados_xlsx: Path
    resultados1_xlsx: Path
    plano_php: Path
    plano_11678: Path
    plano_11679: Path
    afiliados_jotform_csv: Path
    plano_nov_php: Path
    ruta_inclusion_zip: Path
    models_zip: Path


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    raw = _read_text(path)
    reader = csv.DictReader(raw.splitlines())
    rows = []
    for row in reader:
        rows.append({(k or "").strip(): (v or "").strip() for k, v in row.items()})
    return rows


def _inspect_xlsx(path: Path) -> dict[str, Any]:
    if load_workbook is None:
        return {"path": str(path), "ok": False, "error": "openpyxl no disponible"}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover
        return {"path": str(path), "ok": False, "error": str(exc)}

    sheets: list[dict[str, Any]] = []
    for name in wb.sheetnames[:3]:
        ws = wb[name]
        rows = ws.iter_rows(min_row=1, max_row=3, values_only=True)
        data = list(rows)
        header = list(data[0]) if data else []
        preview = list(data[1]) if len(data) > 1 else []
        sheets.append(
            {
                "name": name,
                "header_preview": ["" if v is None else str(v) for v in header[:20]],
                "row2_preview": ["" if v is None else str(v) for v in preview[:20]],
            }
        )
    return {"path": str(path), "ok": True, "sheets": sheets}


def _inspect_plano_txt(path: Path) -> dict[str, Any]:
    lines = _read_text(path).splitlines()
    lengths = [len(line) for line in lines if line]
    type_counts: Counter[str] = Counter()
    for line in lines:
        if not line:
            continue
        type_counts[line[:1]] += 1
    return {
        "path": str(path),
        "line_count": len(lines),
        "record_type_counts": dict(type_counts),
        "line_lengths": sorted(set(lengths)),
    }


def _inspect_php(path: Path) -> dict[str, Any]:
    content = _read_text(path)
    marker_counts = {
        "function tipoPersona": content.count("function tipoPersona"),
        "function traesubtipoafiliado": content.count("function traesubtipoafiliado"),
        "function validacionTrabajadorTemporal": content.count("function validacionTrabajadorTemporal"),
        "linea1_builder": content.count("$linea1 ="),
        "linea3_builder": content.count("$linea3 ="),
        "update_plano": content.count("set plano = '1'"),
    }
    return {"path": str(path), "markers": marker_counts, "size_bytes": path.stat().st_size}


def _inspect_zip(path: Path, max_files: int = 120) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            ext_counter: Counter[str] = Counter()
            for name in names:
                p = Path(name)
                ext = p.suffix.lower() or "[sin_extension]"
                ext_counter[ext] += 1
            return {
                "path": str(path),
                "file_count": len(names),
                "extensions": dict(ext_counter.most_common(20)),
                "files_preview": names[:max_files],
            }
    except Exception as exc:  # pragma: no cover
        return {"path": str(path), "error": str(exc)}


def _kpi(rows: list[dict[str, str]]) -> dict[str, Any]:
    estado = Counter((r.get("estado", "").strip() or "SIN_ESTADO") for r in rows)
    tipo_doc = Counter((r.get("tipodocumento", "").strip() or "SIN_TIPO") for r in rows)
    tipo_doc_emp = Counter((r.get("tipodocumentoempleador", "").strip() or "SIN_TIPO") for r in rows)
    return {
        "count": len(rows),
        "estado_top": dict(estado.most_common(10)),
        "tipodocumento_top": dict(tipo_doc.most_common(10)),
        "tipodocumentoempleador_top": dict(tipo_doc_emp.most_common(10)),
    }


def _build_catalog(*datasets: tuple[str, list[dict[str, str]]]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for source, rows in datasets:
        for row in rows:
            out.append(
                {
                    "source": source,
                    "idtramite": row.get("idtramite", ""),
                    "estado": row.get("estado", ""),
                    "tipodocumento": row.get("tipodocumento", ""),
                    "numerodocumento": row.get("numerodocumento", ""),
                    "tipodocumentoempleador": row.get("tipodocumentoempleador", ""),
                    "numerodocumentoempleador": row.get("numerodocumentoempleador", ""),
                    "razonsocialempleador": row.get("razonsocialempleador", ""),
                    "naturalezajuridica": row.get("naturalezajuridica", ""),
                    "tipoafiliacion": row.get("tipoafiliacion", ""),
                    "tipoaportante": row.get("tipoaportante", ""),
                    "iniciocontrato": row.get("iniciocontrato", ""),
                    "finalizacioncontrato": row.get("finalizacioncontrato", ""),
                    "ibc": row.get("ibc", ""),
                    "iniciocobertura": row.get("iniciocobertura", ""),
                }
            )
    return out


def _write_catalog_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def run(inputs: InputSet, out_dir: Path) -> dict[str, Any]:
    afiliados = _read_csv_rows(inputs.afiliados_csv)
    devuelto = _read_csv_rows(inputs.devuelto_csv)
    dewvuelto = _read_csv_rows(inputs.dewvuelto_csv)
    afiliados_jotform = _read_csv_rows(inputs.afiliados_jotform_csv)
    catalog = _build_catalog(
        ("afiliados.csv", afiliados),
        ("devuelto.csv", devuelto),
        ("dewvuelto jotform.csv", dewvuelto),
        ("afiliados jotform.csv", afiliados_jotform),
    )

    by_idtramite = Counter(r.get("idtramite", "") for r in catalog if r.get("idtramite"))
    duplicated_idtramite = [k for k, v in by_idtramite.items() if v > 1][:30]

    summary = {
        "generated_at": datetime.now().isoformat(),
        "inputs": {
            "csv": [
                str(inputs.afiliados_csv),
                str(inputs.devuelto_csv),
                str(inputs.dewvuelto_csv),
                str(inputs.afiliados_jotform_csv),
            ],
            "xlsx": [str(inputs.informe_xlsx), str(inputs.resultados_xlsx), str(inputs.resultados1_xlsx)],
            "legacy_code": [str(inputs.plano_php), str(inputs.plano_nov_php)],
            "legacy_plans": [str(inputs.plano_11678), str(inputs.plano_11679)],
            "legacy_zips": [str(inputs.ruta_inclusion_zip), str(inputs.models_zip)],
        },
        "csv_kpi": {
            "afiliados": _kpi(afiliados),
            "devuelto": _kpi(devuelto),
            "dewvuelto_jotform": _kpi(dewvuelto),
            "afiliados_jotform": _kpi(afiliados_jotform),
            "catalog_total": len(catalog),
            "duplicated_idtramite_preview": duplicated_idtramite,
        },
        "xlsx_probe": {
            "informe_tramites": _inspect_xlsx(inputs.informe_xlsx),
            "resultados_exec_ruta": _inspect_xlsx(inputs.resultados_xlsx),
            "resultados_exec_ruta_1": _inspect_xlsx(inputs.resultados1_xlsx),
        },
        "legacy_probe": {
            "plano_php": _inspect_php(inputs.plano_php),
            "plano_nov_php": _inspect_php(inputs.plano_nov_php),
            "plano_11678": _inspect_plano_txt(inputs.plano_11678),
            "plano_11679": _inspect_plano_txt(inputs.plano_11679),
            "ruta_inclusion_zip": _inspect_zip(inputs.ruta_inclusion_zip),
            "models_zip": _inspect_zip(inputs.models_zip),
        },
        "target_system_bootstrap": {
            "phase_1": "Ingestion y normalizacion de fuentes legacy",
            "phase_2": "Motor de reglas (validaciones tipo documento, tipo persona y mapeos de codigos)",
            "phase_3": "Generador de plano fijo + validador de longitud por registro",
            "phase_4": "Reproceso, devoluciones y trazabilidad por idtramite/lote",
        },
    }

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir.mkdir(parents=True, exist_ok=True)
    summary_path = out_dir / f"legacy_bootstrap_summary_{ts}.json"
    catalog_path = out_dir / f"legacy_catalog_{ts}.csv"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_catalog_csv(catalog_path, catalog)
    summary["outputs"] = {"summary_json": str(summary_path), "catalog_csv": str(catalog_path)}
    return summary


def main() -> None:
    default_input_dir = Path(os.getenv("AFILEGA_LEGACY_INPUT_DIR", str(Path(__file__).resolve().parents[2] / "shared_downloads" / "legacy")))
    parser = argparse.ArgumentParser(description="Bootstrap de nuevo sistema de afiliaciones desde insumos legacy.")
    parser.add_argument(
        "--out-dir",
        default=str(Path(__file__).resolve().parents[2] / "docs"),
        help="Carpeta de salida para resumen y catalogo.",
    )
    parser.add_argument("--afiliados-csv", default=str(default_input_dir / "afiliados.csv"))
    parser.add_argument("--devuelto-csv", default=str(default_input_dir / "devuelto.csv"))
    parser.add_argument("--dewvuelto-csv", default=str(default_input_dir / "dewvuelto jotform.csv"))
    parser.add_argument("--informe-xlsx", default=str(default_input_dir / "informe_tramites.xlsx"))
    parser.add_argument("--resultados-xlsx", default=str(default_input_dir / "resultados EXEC ruta.xlsx"))
    parser.add_argument("--resultados1-xlsx", default=str(default_input_dir / "resultados EXEC ruta-1.xlsx"))
    parser.add_argument("--plano-php", default=str(default_input_dir / "planoAfi1.php"))
    parser.add_argument("--plano-11678", default=str(default_input_dir / "planoIndp_11678.txt"))
    parser.add_argument("--plano-11679", default=str(default_input_dir / "planoIndp_11679.txt"))
    parser.add_argument("--afiliados-jotform-csv", default=str(default_input_dir / "afiliados jotform.csv"))
    parser.add_argument("--plano-nov-php", default=str(default_input_dir / "planoAfiNov.php"))
    parser.add_argument("--ruta-inclusion-zip", default=str(default_input_dir / "ruta_inclusion.zip"))
    parser.add_argument("--models-zip", default=str(default_input_dir / "models.zip"))
    args = parser.parse_args()

    inputs = InputSet(
        afiliados_csv=Path(args.afiliados_csv),
        devuelto_csv=Path(args.devuelto_csv),
        dewvuelto_csv=Path(args.dewvuelto_csv),
        informe_xlsx=Path(args.informe_xlsx),
        resultados_xlsx=Path(args.resultados_xlsx),
        resultados1_xlsx=Path(args.resultados1_xlsx),
        plano_php=Path(args.plano_php),
        plano_11678=Path(args.plano_11678),
        plano_11679=Path(args.plano_11679),
        afiliados_jotform_csv=Path(args.afiliados_jotform_csv),
        plano_nov_php=Path(args.plano_nov_php),
        ruta_inclusion_zip=Path(args.ruta_inclusion_zip),
        models_zip=Path(args.models_zip),
    )
    out = run(inputs, Path(args.out_dir))
    print(json.dumps(out["outputs"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
